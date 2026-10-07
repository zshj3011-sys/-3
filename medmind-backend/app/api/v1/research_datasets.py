"""科研数据导入 API — PRD §07 用户故事 #165 (Must · V1.0)

接受 .csv / .xlsx / .sav 上传, 嗅探列类型, 生成预览, 持久化元数据.

设计选择 (v3.6.2):
- CSV: Python stdlib csv (零依赖)
- Excel xlsx: openpyxl (已加入 requirements)
- SPSS .sav: pyreadstat 可选依赖. 未安装时返回 501 并提示客户安装.
- 文件本体 mock 模式不写盘, 仅保留 preview 与 schema 元数据;
  生产模式应注入对象存储客户端, 写 content_key.
"""
import csv as _csv
import io
import json
from datetime import datetime
from typing import Optional, List, Any
from fastapi import APIRouter, Depends, HTTPException, UploadFile, File, Form, Query
from pydantic import BaseModel, Field, ConfigDict
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy import select, func, desc

from app.core.database import get_db
from app.core.deps import get_current_user
from app.models.research_dataset import ResearchDataset
from app.models.user import User
from app.schemas.common import StandardResponse, PageResponse, PageMeta, OrmBase

router = APIRouter()

ALLOWED_EXTS = {"csv", "xlsx", "sav"}
MAX_BYTES = 20 * 1024 * 1024  # 20MB 上限 (演示)
PREVIEW_ROWS = 50


# ============== Schema ==============
class DatasetOut(OrmBase):
    id: int
    name: str
    original_filename: str
    file_format: str
    size_bytes: int
    rows: int
    columns: int
    notes: Optional[str]
    created_at: datetime


class DatasetPreview(BaseModel):
    # 避免 Pydantic 与 BaseModel.schema() 冲突
    model_config = ConfigDict(protected_namespaces=())

    id: int
    name: str
    file_format: str
    rows: int
    columns: int
    schema_: List[dict] = Field(..., alias="schema",
                                description="列元数据: name/dtype/null_count/sample")
    preview: List[dict] = Field(..., description="前 N 行数据")


# ============== Helpers: 嗅探列类型 ==============
def _infer_dtype(values: List[Any]) -> str:
    """简单类型嗅探: int / float / bool / date / str"""
    non_null = [v for v in values if v not in (None, "", "NA", "NaN")]
    if not non_null:
        return "unknown"
    # 全部是 bool
    if all(isinstance(v, bool) for v in non_null):
        return "bool"
    # int
    try:
        for v in non_null:
            int(str(v))
        return "int"
    except (ValueError, TypeError):
        pass
    # float
    try:
        for v in non_null:
            float(str(v))
        return "float"
    except (ValueError, TypeError):
        pass
    # date
    for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%Y-%m-%d %H:%M:%S"):
        try:
            for v in non_null[:5]:
                datetime.strptime(str(v), fmt)
            return "date"
        except (ValueError, TypeError):
            continue
    return "str"


def _parse_csv(content: bytes) -> tuple[List[str], List[List[Any]]]:
    text = content.decode("utf-8-sig", errors="replace")
    reader = _csv.reader(io.StringIO(text))
    rows = list(reader)
    if not rows:
        return [], []
    header = [h.strip() for h in rows[0]]
    body = rows[1:]
    return header, body


def _parse_xlsx(content: bytes) -> tuple[List[str], List[List[Any]]]:
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise HTTPException(501, "服务器未安装 openpyxl 依赖, 无法解析 xlsx")
    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return [], []
    header = [str(h) if h is not None else "" for h in rows[0]]
    body = [list(r) for r in rows[1:]]
    return header, body


def _parse_sav(content: bytes) -> tuple[List[str], List[List[Any]]]:
    """SPSS .sav 解析 — 可选依赖"""
    try:
        import pyreadstat  # type: ignore
    except ImportError:
        raise HTTPException(
            501,
            "SPSS .sav 支持需要安装可选依赖 pyreadstat. "
            "运行: pip install pyreadstat>=1.2.7 后重启服务. "
            "客户上线检查清单已记录此项."
        )
    # pyreadstat 需要文件路径, 此处写临时文件
    import tempfile, os
    tmp = tempfile.NamedTemporaryFile(suffix=".sav", delete=False)
    try:
        tmp.write(content)
        tmp.close()
        df, meta = pyreadstat.read_sav(tmp.name)
        header = list(df.columns)
        body = df.values.tolist()
        return header, body
    finally:
        try:
            os.unlink(tmp.name)
        except OSError:
            pass


def _build_schema_and_preview(header: List[str], body: List[List[Any]]):
    """根据 header+body 推断 schema 与前 N 行预览."""
    schema = []
    for col_idx, col_name in enumerate(header):
        col_vals = [row[col_idx] if col_idx < len(row) else None for row in body]
        dtype = _infer_dtype(col_vals)
        null_count = sum(1 for v in col_vals if v in (None, "", "NA", "NaN"))
        sample = next((v for v in col_vals if v not in (None, "")), None)
        schema.append({
            "name": col_name or f"col_{col_idx+1}",
            "dtype": dtype,
            "null_count": null_count,
            "sample": str(sample) if sample is not None else None,
        })
    preview = []
    for row in body[:PREVIEW_ROWS]:
        d = {}
        for col_idx, col_name in enumerate(header):
            v = row[col_idx] if col_idx < len(row) else None
            # 保证 JSON 可序列化
            if isinstance(v, datetime):
                v = v.isoformat()
            d[col_name or f"col_{col_idx+1}"] = v
        preview.append(d)
    return schema, preview


# ============== 端点 ==============
@router.post("/upload", response_model=StandardResponse[DatasetPreview], summary="上传科研数据集")
async def upload_dataset(
    file: UploadFile = File(...),
    name: Optional[str] = Form(None, description="可选: 自定义数据集名称"),
    notes: Optional[str] = Form(None),
    db: AsyncSession = Depends(get_db),
    current: User = Depends(get_current_user),
):
    """PRD §07 用户故事 #165: 研究者导入 Excel/CSV/SPSS 数据.

    返回数据集元数据 + 列 schema + 前 50 行预览, 后续可用于 statistics/recommend.
    """
    filename = file.filename or "unknown"
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if ext not in ALLOWED_EXTS:
        raise HTTPException(400, f"不支持的格式 .{ext}; 支持: {sorted(ALLOWED_EXTS)}")

    content = await file.read()
    if len(content) > MAX_BYTES:
        raise HTTPException(413, f"文件超过 {MAX_BYTES//1024//1024}MB 上限")
    if len(content) == 0:
        raise HTTPException(400, "文件为空")

    # 解析
    if ext == "csv":
        header, body = _parse_csv(content)
    elif ext == "xlsx":
        header, body = _parse_xlsx(content)
    elif ext == "sav":
        header, body = _parse_sav(content)
    else:
        raise HTTPException(400, "不支持的格式")

    if not header:
        raise HTTPException(400, "文件中没有可识别的列头")

    schema, preview = _build_schema_and_preview(header, body)

    ds = ResearchDataset(
        owner_id=current.id,
        name=name or filename.rsplit(".", 1)[0],
        original_filename=filename,
        file_format=ext,
        size_bytes=len(content),
        rows=len(body),
        columns=len(header),
        schema_json=schema,
        preview_json=preview,
        notes=notes,
    )
    db.add(ds)
    await db.commit()
    await db.refresh(ds)

    return StandardResponse(data=DatasetPreview.model_validate({
        "id": ds.id, "name": ds.name, "file_format": ds.file_format,
        "rows": ds.rows, "columns": ds.columns,
        "schema": schema, "preview": preview,
    }))


@router.get("/", response_model=PageResponse[DatasetOut], summary="我的数据集列表")
async def list_datasets(
    page: int = Query(1, ge=1),
    page_size: int = Query(20, ge=1, le=100),
    db: AsyncSession = Depends(get_db),
    current: User = Depends(get_current_user),
):
    stmt = select(ResearchDataset).where(ResearchDataset.owner_id == current.id)\
        .order_by(desc(ResearchDataset.created_at))
    total = (await db.execute(
        select(func.count(ResearchDataset.id)).where(ResearchDataset.owner_id == current.id)
    )).scalar() or 0
    stmt = stmt.offset((page-1)*page_size).limit(page_size)
    rows = (await db.execute(stmt)).scalars().all()
    data = [DatasetOut.model_validate(r) for r in rows]
    meta = PageMeta(page=page, page_size=page_size, total=total,
                    total_pages=(total + page_size - 1) // page_size if total else 0)
    return PageResponse(data=data, meta=meta)


@router.get("/{dataset_id}/preview", response_model=StandardResponse[DatasetPreview],
            summary="数据集预览 (schema + 前 50 行)")
async def preview_dataset(
    dataset_id: int,
    db: AsyncSession = Depends(get_db),
    current: User = Depends(get_current_user),
):
    ds = await db.get(ResearchDataset, dataset_id)
    if not ds:
        raise HTTPException(404, "数据集不存在")
    if ds.owner_id != current.id and current.role != "admin":
        raise HTTPException(403, "无权访问此数据集")
    return StandardResponse(data=DatasetPreview.model_validate({
        "id": ds.id, "name": ds.name, "file_format": ds.file_format,
        "rows": ds.rows, "columns": ds.columns,
        "schema": ds.schema_json or [],
        "preview": ds.preview_json or [],
    }))


@router.delete("/{dataset_id}", response_model=StandardResponse[dict], summary="删除数据集")
async def delete_dataset(
    dataset_id: int,
    db: AsyncSession = Depends(get_db),
    current: User = Depends(get_current_user),
):
    ds = await db.get(ResearchDataset, dataset_id)
    if not ds:
        raise HTTPException(404, "数据集不存在")
    if ds.owner_id != current.id and current.role != "admin":
        raise HTTPException(403, "无权删除此数据集")
    await db.delete(ds)
    await db.commit()
    return StandardResponse(data={"deleted": dataset_id})
